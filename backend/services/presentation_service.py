"""
Presentation Service - Generates PowerPoint presentations from research reports.

This service takes research content and:
1. Uses an LLM to structure the content into 10 presentation slides
2. Generates a professional PPTX file using python-pptx

Supports multiple LLM providers for slide structuring.
"""

import os
import io
import json
import logging
import re
from typing import Dict, List, Optional, Literal
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from dotenv import load_dotenv

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# Load .env from project root
env_path = Path(__file__).parent.parent.parent / ".env"
load_dotenv(env_path)

logger = logging.getLogger(__name__)


@dataclass
class SlideData:
    """Data for a single slide."""
    slide_number: int
    title: str
    content_type: str  # "title", "bullets", "table", "text"
    content: List[str] | List[List[str]]  # Bullets, table rows, or text paragraphs
    speaker_notes: str = ""


@dataclass
class PresentationData:
    """Complete presentation data from LLM."""
    slides: List[SlideData]
    company: str
    prompt_name: str
    generated_at: datetime


# Prompt template for structuring content into slides
SLIDE_STRUCTURE_PROMPT = """You are an expert presentation designer. Analyze this competitive intelligence research report and structure it into exactly 10 professional presentation slides.

**COMPANY:** {company}
**REPORT TYPE:** {prompt_name}

## SLIDE STRUCTURE REQUIREMENTS

Create exactly 10 slides with this structure:
1. **Title Slide** - Company name, report title, date, "CONFIDENTIAL"
2. **Executive Summary** - 4-5 key bullet points summarizing the most important findings
3. **Key Findings (Part 1)** - First set of critical findings (4-6 bullets)
4. **Key Findings (Part 2)** - Second set of critical findings (4-6 bullets)
5. **Critical Intelligence** - Most urgent/important intelligence items (4-5 bullets)
6. **Market Position** - Competitive positioning, market trends (4-6 bullets)
7. **Data & Metrics** - Key numbers, statistics as a table (headers + 4-6 data rows)
8. **Strategic Implications** - What this means for decision-makers (4-5 bullets)
9. **Recommendations** - Actionable next steps (4-5 bullets)
10. **Sources & Methodology** - List of sources cited, methodology note (4-6 bullets)

## OUTPUT FORMAT

Output a JSON array with exactly 10 objects. Each object must have:
- "slide_number": 1-10
- "title": Clear, concise slide title (max 60 chars)
- "content_type": One of "title", "bullets", "table", "text"
- "content": Array of content items:
  - For "title": ["subtitle text"]
  - For "bullets": ["bullet 1", "bullet 2", ...]
  - For "table": [["Header1", "Header2", ...], ["row1col1", "row1col2", ...], ...]
  - For "text": ["paragraph 1", "paragraph 2"]
- "speaker_notes": DETAILED presenter notes (see Speaker Notes Requirements below)

## SPEAKER NOTES REQUIREMENTS

Speaker notes are CRITICAL. Write them as a comprehensive briefing paragraph that helps the presenter deeply understand the slide content.

DO NOT just repeat the bullet points. Instead, write speaker notes that:
- Explain the "so what" - why this information matters
- Provide background context and details from the research that didn't fit on the slide
- Connect the dots between different points on the slide
- Include specific examples, numbers, dates, or evidence that support the bullets
- Highlight implications and what decisions or actions should result from this information
- Note any caveats, uncertainties, or areas where sources disagreed

Write speaker notes as flowing prose paragraphs, not a list. They should read like a knowledgeable analyst briefing an executive on what they need to know about this slide. Be generous with detail - the presenter should be able to read the notes and feel fully prepared to discuss the topic in depth.

## GUIDELINES

- Keep bullet points concise (max 100 characters each)
- Use action-oriented language
- Include specific numbers and data where available
- For tables, keep to max 5 columns and 6 rows
- Prioritize actionable intelligence over generic observations

## RESEARCH CONTENT

{content}

## OUTPUT

Output ONLY the JSON array, no additional text or markdown formatting:"""


class PresentationService:
    """
    Service for generating PowerPoint presentations from research content.
    Uses LLM to structure content, then generates PPTX with python-pptx.
    """

    # Color scheme
    PRIMARY_COLOR = RGBColor(0x1e, 0x40, 0xaf)    # Blue-800
    SECONDARY_COLOR = RGBColor(0x3b, 0x82, 0xf6)  # Blue-500
    ACCENT_COLOR = RGBColor(0xdc, 0x26, 0x26)     # Red-600 (Red 6 branding)
    TEXT_COLOR = RGBColor(0x1f, 0x29, 0x37)       # Gray-800
    LIGHT_BG = RGBColor(0xf8, 0xfa, 0xfc)         # Slate-50

    def __init__(self, llm_provider: str = "google"):
        """
        Initialize the presentation service.

        Args:
            llm_provider: Which LLM to use for structuring ("google", "openai", "anthropic", "xai")
        """
        self.llm_provider = llm_provider
        self._setup_llm_client()

    def _setup_llm_client(self):
        """Set up the LLM client based on provider."""
        if self.llm_provider == "google":
            api_key = os.getenv("GEMINI_API_KEY")
            if not api_key:
                raise ValueError("GEMINI_API_KEY not found")
            from google import genai
            self.client = genai.Client(api_key=api_key)
            self.model = "gemini-2.0-flash"

        elif self.llm_provider == "openai":
            api_key = os.getenv("OPENAI_API_KEY")
            if not api_key:
                raise ValueError("OPENAI_API_KEY not found")
            from openai import OpenAI
            self.client = OpenAI(api_key=api_key)
            self.model = "gpt-4o"

        elif self.llm_provider == "anthropic":
            api_key = os.getenv("ANTHROPIC_API_KEY")
            if not api_key:
                raise ValueError("ANTHROPIC_API_KEY not found")
            import anthropic
            self.client = anthropic.Anthropic(api_key=api_key)
            self.model = "claude-sonnet-4-20250514"

        elif self.llm_provider == "xai":
            api_key = os.getenv("XAI_API_KEY")
            if not api_key:
                raise ValueError("XAI_API_KEY not found")
            from openai import OpenAI
            self.client = OpenAI(api_key=api_key, base_url="https://api.x.ai/v1")
            self.model = "grok-3-latest"

        else:
            raise ValueError(f"Unsupported LLM provider: {self.llm_provider}")

    def generate_presentation(
        self,
        content: str,
        company: str,
        prompt_name: str,
        citations: List[Dict] = None
    ) -> bytes:
        """
        Generate a PowerPoint presentation from research content.

        Args:
            content: Research report content (markdown)
            company: Company name
            prompt_name: Research prompt/report type name
            citations: Optional list of citation dicts with url and title

        Returns:
            PPTX file as bytes
        """
        # Step 1: Structure content into slides using LLM
        slide_data = self._structure_content_for_slides(content, company, prompt_name)

        # Step 2: Add citations to the sources slide if available
        if citations and len(slide_data) >= 10:
            sources_slide = slide_data[9]  # Slide 10 (0-indexed)
            citation_bullets = [f"{c.get('title', 'Source')}: {c.get('url', '')}" for c in citations[:6]]
            if citation_bullets:
                sources_slide.content = citation_bullets + sources_slide.content[:2]

        # Step 3: Generate PPTX file
        pptx_bytes = self._create_pptx(slide_data, company, prompt_name)

        return pptx_bytes

    def generate_slide_text(
        self,
        content: str,
        company: str,
        prompt_name: str,
        citations: List[Dict] = None
    ) -> str:
        """
        Generate formatted text for 10 slides that can be copied into a slide deck.

        Args:
            content: Research report content (markdown)
            company: Company name
            prompt_name: Research prompt/report type name
            citations: Optional list of citation dicts with url and title

        Returns:
            Formatted text with 10 slide sections
        """
        # Step 1: Structure content into slides using LLM
        slide_data = self._structure_content_for_slides(content, company, prompt_name)

        # Step 2: Add citations to the sources slide if available
        if citations and len(slide_data) >= 10:
            sources_slide = slide_data[9]  # Slide 10 (0-indexed)
            citation_bullets = [f"{c.get('title', 'Source')}: {c.get('url', '')}" for c in citations[:6]]
            if citation_bullets:
                sources_slide.content = citation_bullets + sources_slide.content[:2]

        # Step 3: Format as text
        return self._format_slides_as_text(slide_data, company, prompt_name)

    def _format_slides_as_text(
        self,
        slides: List[SlideData],
        company: str,
        prompt_name: str
    ) -> str:
        """Format slide data as copyable text sections."""
        lines = []
        lines.append("=" * 80)
        lines.append(f"SLIDE DECK CONTENT: {company} - {prompt_name}")
        lines.append(f"Generated: {datetime.now().strftime('%B %d, %Y')}")
        lines.append("=" * 80)
        lines.append("")
        lines.append("Instructions: Copy each slide section below into your branded template.")
        lines.append("Each section is formatted for easy copy/paste.")
        lines.append("")

        for slide in slides:
            lines.append("-" * 80)
            lines.append(f"SLIDE {slide.slide_number}: {slide.title}")
            lines.append("-" * 80)
            lines.append("")

            if slide.content_type == "title":
                # Title slide - just show the subtitle/content
                for item in slide.content:
                    lines.append(f"  {item}")
            elif slide.content_type == "table":
                # Table - format as aligned columns
                if slide.content and isinstance(slide.content[0], list):
                    for row in slide.content:
                        lines.append("  | " + " | ".join(str(cell) for cell in row) + " |")
                else:
                    for item in slide.content:
                        lines.append(f"  • {item}")
            else:
                # Bullets or text
                for item in slide.content:
                    text = str(item).strip()
                    if text.startswith("- ") or text.startswith("• "):
                        text = text[2:]
                    lines.append(f"  • {text}")

            lines.append("")

            if slide.speaker_notes:
                lines.append(f"  [Speaker Notes: {slide.speaker_notes}]")
                lines.append("")

        lines.append("=" * 80)
        lines.append("END OF SLIDE CONTENT")
        lines.append("=" * 80)

        return "\n".join(lines)

    def _structure_content_for_slides(
        self,
        content: str,
        company: str,
        prompt_name: str
    ) -> List[SlideData]:
        """
        Use LLM to structure content into 10 slides.

        Args:
            content: Research content (markdown)
            company: Company name
            prompt_name: Report type name

        Returns:
            List of SlideData objects
        """
        prompt = SLIDE_STRUCTURE_PROMPT.format(
            company=company,
            prompt_name=prompt_name,
            content=content[:15000]  # Truncate if very long
        )

        try:
            # Call the appropriate LLM
            if self.llm_provider == "google":
                response = self.client.models.generate_content(
                    model=self.model,
                    contents=prompt
                )
                response_text = response.text

            elif self.llm_provider in ["openai", "xai"]:
                response = self.client.chat.completions.create(
                    model=self.model,
                    messages=[{"role": "user", "content": prompt}],
                    temperature=0.3
                )
                response_text = response.choices[0].message.content

            elif self.llm_provider == "anthropic":
                response = self.client.messages.create(
                    model=self.model,
                    max_tokens=4096,
                    messages=[{"role": "user", "content": prompt}]
                )
                response_text = response.content[0].text

            # Parse JSON response
            slides = self._parse_slide_json(response_text, company, prompt_name)
            return slides

        except Exception as e:
            logger.error(f"LLM slide structuring failed: {e}")
            # Return default structure on failure
            return self._get_default_slides(company, prompt_name, content)

    def _parse_slide_json(
        self,
        response_text: str,
        company: str,
        prompt_name: str
    ) -> List[SlideData]:
        """Parse LLM JSON response into SlideData objects."""
        try:
            # Try to extract JSON from response
            json_match = re.search(r'\[[\s\S]*\]', response_text)
            if json_match:
                json_str = json_match.group()
                slides_data = json.loads(json_str)
            else:
                raise ValueError("No JSON array found in response")

            slides = []
            for slide_dict in slides_data:
                slides.append(SlideData(
                    slide_number=slide_dict.get("slide_number", len(slides) + 1),
                    title=slide_dict.get("title", f"Slide {len(slides) + 1}"),
                    content_type=slide_dict.get("content_type", "bullets"),
                    content=slide_dict.get("content", []),
                    speaker_notes=slide_dict.get("speaker_notes", "")
                ))

            # Ensure we have exactly 10 slides
            while len(slides) < 10:
                slides.append(self._get_placeholder_slide(len(slides) + 1))

            return slides[:10]

        except (json.JSONDecodeError, ValueError) as e:
            logger.error(f"Failed to parse slide JSON: {e}")
            return self._get_default_slides(company, prompt_name, "")

    def _get_placeholder_slide(self, num: int) -> SlideData:
        """Get a placeholder slide."""
        return SlideData(
            slide_number=num,
            title=f"Additional Content {num}",
            content_type="bullets",
            content=["Content to be added"],
            speaker_notes=""
        )

    def _get_default_slides(
        self,
        company: str,
        prompt_name: str,
        content: str
    ) -> List[SlideData]:
        """Get default slide structure when LLM fails."""
        return [
            SlideData(1, f"{company}", "title", [prompt_name, datetime.now().strftime("%B %d, %Y")], ""),
            SlideData(2, "Executive Summary", "bullets", ["Key findings from research", "Strategic implications", "Recommended actions"], ""),
            SlideData(3, "Key Findings", "bullets", ["Finding 1", "Finding 2", "Finding 3", "Finding 4"], ""),
            SlideData(4, "Key Findings (Continued)", "bullets", ["Finding 5", "Finding 6", "Finding 7"], ""),
            SlideData(5, "Critical Intelligence", "bullets", ["Critical item 1", "Critical item 2", "Critical item 3"], ""),
            SlideData(6, "Market Position", "bullets", ["Market position insight 1", "Market position insight 2"], ""),
            SlideData(7, "Data & Metrics", "bullets", ["Metric 1", "Metric 2", "Metric 3"], ""),
            SlideData(8, "Strategic Implications", "bullets", ["Implication 1", "Implication 2"], ""),
            SlideData(9, "Recommendations", "bullets", ["Recommendation 1", "Recommendation 2", "Recommendation 3"], ""),
            SlideData(10, "Sources & Methodology", "bullets", ["AI-powered competitive intelligence analysis", f"Report generated: {datetime.now().strftime('%Y-%m-%d')}"], ""),
        ]

    def _create_pptx(
        self,
        slides: List[SlideData],
        company: str,
        prompt_name: str
    ) -> bytes:
        """
        Create PPTX file from slide data.

        Args:
            slides: List of SlideData objects
            company: Company name
            prompt_name: Report type

        Returns:
            PPTX file as bytes
        """
        prs = Presentation()
        prs.slide_width = Inches(13.333)  # 16:9 aspect ratio
        prs.slide_height = Inches(7.5)

        for slide_data in slides:
            if slide_data.content_type == "title":
                self._add_title_slide(prs, slide_data, company)
            elif slide_data.content_type == "table":
                self._add_table_slide(prs, slide_data)
            else:  # bullets or text
                self._add_content_slide(prs, slide_data)

        # Save to bytes
        pptx_buffer = io.BytesIO()
        prs.save(pptx_buffer)
        pptx_buffer.seek(0)
        return pptx_buffer.getvalue()

    def _add_title_slide(self, prs: Presentation, slide_data: SlideData, company: str):
        """Add a title slide."""
        slide_layout = prs.slide_layouts[6]  # Blank layout
        slide = prs.slides.add_slide(slide_layout)

        # Add blue header bar
        header = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0), Inches(0),
            prs.slide_width, Inches(2.5)
        )
        header.fill.solid()
        header.fill.fore_color.rgb = self.PRIMARY_COLOR
        header.line.fill.background()

        # Company name (main title)
        title_box = slide.shapes.add_textbox(
            Inches(0.5), Inches(0.8), Inches(12), Inches(1)
        )
        tf = title_box.text_frame
        p = tf.paragraphs[0]
        p.text = company
        p.font.size = Pt(44)
        p.font.bold = True
        p.font.color.rgb = RGBColor(0xff, 0xff, 0xff)
        p.alignment = PP_ALIGN.CENTER

        # Subtitle (report type)
        subtitle_box = slide.shapes.add_textbox(
            Inches(0.5), Inches(1.8), Inches(12), Inches(0.5)
        )
        tf = subtitle_box.text_frame
        p = tf.paragraphs[0]
        subtitle_text = slide_data.content[0] if slide_data.content else prompt_name
        p.text = subtitle_text
        p.font.size = Pt(24)
        p.font.color.rgb = RGBColor(0xff, 0xff, 0xff)
        p.alignment = PP_ALIGN.CENTER

        # Date
        date_box = slide.shapes.add_textbox(
            Inches(0.5), Inches(4), Inches(12), Inches(0.5)
        )
        tf = date_box.text_frame
        p = tf.paragraphs[0]
        date_text = slide_data.content[1] if len(slide_data.content) > 1 else datetime.now().strftime("%B %d, %Y")
        p.text = date_text
        p.font.size = Pt(18)
        p.font.color.rgb = self.TEXT_COLOR
        p.alignment = PP_ALIGN.CENTER

        # Confidential marker
        conf_box = slide.shapes.add_textbox(
            Inches(0.5), Inches(6.5), Inches(12), Inches(0.4)
        )
        tf = conf_box.text_frame
        p = tf.paragraphs[0]
        p.text = "CONFIDENTIAL - RED 6 COMPETITIVE INTELLIGENCE"
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = self.ACCENT_COLOR
        p.alignment = PP_ALIGN.CENTER

    def _add_content_slide(self, prs: Presentation, slide_data: SlideData):
        """Add a content slide with bullets or text."""
        slide_layout = prs.slide_layouts[6]  # Blank layout
        slide = prs.slides.add_slide(slide_layout)

        # Title bar
        title_bar = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0), Inches(0),
            prs.slide_width, Inches(1.2)
        )
        title_bar.fill.solid()
        title_bar.fill.fore_color.rgb = self.PRIMARY_COLOR
        title_bar.line.fill.background()

        # Title text
        title_box = slide.shapes.add_textbox(
            Inches(0.5), Inches(0.35), Inches(12), Inches(0.6)
        )
        tf = title_box.text_frame
        p = tf.paragraphs[0]
        p.text = slide_data.title
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = RGBColor(0xff, 0xff, 0xff)

        # Content area
        content_box = slide.shapes.add_textbox(
            Inches(0.75), Inches(1.5), Inches(11.5), Inches(5.5)
        )
        tf = content_box.text_frame
        tf.word_wrap = True

        # Add content items
        for i, item in enumerate(slide_data.content):
            if i == 0:
                p = tf.paragraphs[0]
            else:
                p = tf.add_paragraph()

            # Clean the text
            text = str(item).strip()
            if text.startswith("- ") or text.startswith("• "):
                text = text[2:]

            p.text = f"• {text}"
            p.font.size = Pt(18)
            p.font.color.rgb = self.TEXT_COLOR
            p.space_after = Pt(12)
            p.level = 0

        # Slide number
        num_box = slide.shapes.add_textbox(
            Inches(12.5), Inches(7), Inches(0.5), Inches(0.3)
        )
        tf = num_box.text_frame
        p = tf.paragraphs[0]
        p.text = str(slide_data.slide_number)
        p.font.size = Pt(10)
        p.font.color.rgb = RGBColor(0x9c, 0xa3, 0xaf)  # Gray-400
        p.alignment = PP_ALIGN.RIGHT

        # Add speaker notes if present
        if slide_data.speaker_notes:
            notes_slide = slide.notes_slide
            notes_slide.notes_text_frame.text = slide_data.speaker_notes

    def _add_table_slide(self, prs: Presentation, slide_data: SlideData):
        """Add a slide with a table."""
        slide_layout = prs.slide_layouts[6]  # Blank layout
        slide = prs.slides.add_slide(slide_layout)

        # Title bar
        title_bar = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0), Inches(0),
            prs.slide_width, Inches(1.2)
        )
        title_bar.fill.solid()
        title_bar.fill.fore_color.rgb = self.PRIMARY_COLOR
        title_bar.line.fill.background()

        # Title text
        title_box = slide.shapes.add_textbox(
            Inches(0.5), Inches(0.35), Inches(12), Inches(0.6)
        )
        tf = title_box.text_frame
        p = tf.paragraphs[0]
        p.text = slide_data.title
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = RGBColor(0xff, 0xff, 0xff)

        # Parse table data
        table_data = slide_data.content
        if not table_data or not isinstance(table_data[0], list):
            # Fall back to bullets if not proper table format
            self._add_content_slide(prs, slide_data)
            return

        rows = len(table_data)
        cols = len(table_data[0]) if table_data else 1

        # Add table
        table = slide.shapes.add_table(
            rows, cols,
            Inches(0.75), Inches(1.5),
            Inches(11.5), Inches(min(rows * 0.6, 5))
        ).table

        # Style header row
        for j, cell_text in enumerate(table_data[0]):
            cell = table.cell(0, j)
            cell.text = str(cell_text)
            cell.fill.solid()
            cell.fill.fore_color.rgb = self.PRIMARY_COLOR
            p = cell.text_frame.paragraphs[0]
            p.font.color.rgb = RGBColor(0xff, 0xff, 0xff)
            p.font.bold = True
            p.font.size = Pt(14)

        # Style data rows
        for i, row_data in enumerate(table_data[1:], start=1):
            for j, cell_text in enumerate(row_data):
                if j < cols:
                    cell = table.cell(i, j)
                    cell.text = str(cell_text)
                    # Alternating row colors
                    if i % 2 == 0:
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = self.LIGHT_BG
                    p = cell.text_frame.paragraphs[0]
                    p.font.size = Pt(12)
                    p.font.color.rgb = self.TEXT_COLOR

        # Slide number
        num_box = slide.shapes.add_textbox(
            Inches(12.5), Inches(7), Inches(0.5), Inches(0.3)
        )
        tf = num_box.text_frame
        p = tf.paragraphs[0]
        p.text = str(slide_data.slide_number)
        p.font.size = Pt(10)
        p.font.color.rgb = RGBColor(0x9c, 0xa3, 0xaf)
        p.alignment = PP_ALIGN.RIGHT


# Singleton instance factory
_presentation_service: Optional[PresentationService] = None


def get_presentation_service(llm_provider: str = "google") -> PresentationService:
    """Get or create presentation service instance."""
    global _presentation_service
    if _presentation_service is None or _presentation_service.llm_provider != llm_provider:
        _presentation_service = PresentationService(llm_provider)
    return _presentation_service
