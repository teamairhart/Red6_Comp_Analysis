"""Research API Router"""
from fastapi import APIRouter, HTTPException, WebSocket, WebSocketDisconnect, BackgroundTasks
from fastapi.responses import Response
from models.research import (
    ResearchRequest,
    ResearchResponse,
    ResearchJob,
    ResearchStatus,
    ResearchProgress,
    ProviderResult,
    SynthesizedReport,
)
from services.research_service import get_research_service
from services.source_service import (
    get_prompt_keywords,
    find_relevant_sources,
    format_sources_for_prompt,
    match_citations_to_sources,
    update_source_hit_counts,
)
from services.delta_service import run_delta_analysis, run_knowledge_base_delta, save_delta_report
from services.synthesis_service import get_synthesis_service
from services.presentation_service import get_presentation_service
from datetime import datetime
import uuid
import asyncio
from concurrent.futures import ThreadPoolExecutor
from typing import Dict, Literal
import logging
import traceback
import threading
import markdown
import io
import re
import json
from pathlib import Path

# Try to import LangGraph orchestrator (optional dependency)
try:
    from services.langgraph_research_service import get_langgraph_orchestrator
    LANGGRAPH_AVAILABLE = True
except ImportError:
    LANGGRAPH_AVAILABLE = False
    get_langgraph_orchestrator = None

# Try to import Open Deep Research service (for enhanced deep research)
try:
    from services.open_deep_research_service import (
        get_open_deep_research_service,
        run_open_deep_research,
        DeepResearchResult,
    )
    OPEN_DEEP_RESEARCH_AVAILABLE = True
except ImportError:
    OPEN_DEEP_RESEARCH_AVAILABLE = False
    get_open_deep_research_service = None
    run_open_deep_research = None

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/api/research", tags=["research"])

# Storage directory for persisting jobs
JOBS_DIR = Path(__file__).parent.parent / "data" / "jobs"
JOBS_DIR.mkdir(parents=True, exist_ok=True)

# In-memory job storage (loaded from disk on startup)
jobs: Dict[str, ResearchJob] = {}


def _save_job_to_disk(job: ResearchJob):
    """Save a job to disk for persistence."""
    try:
        job_file = JOBS_DIR / f"{job.id}.json"
        job_data = job.model_dump(mode='json')
        # Convert datetime objects to ISO strings
        for key in ['created_at', 'updated_at', 'completed_at']:
            if job_data.get(key):
                if isinstance(job_data[key], datetime):
                    job_data[key] = job_data[key].isoformat()
        job_file.write_text(json.dumps(job_data, indent=2, default=str))
        logger.debug(f"Saved job {job.id} to disk")
    except Exception as e:
        logger.error(f"Failed to save job {job.id} to disk: {e}")


def _load_jobs_from_disk():
    """Load all jobs from disk on startup."""
    global jobs
    loaded_count = 0
    for job_file in JOBS_DIR.glob("*.json"):
        try:
            job_data = json.loads(job_file.read_text())
            # Convert ISO strings back to datetime
            for key in ['created_at', 'updated_at', 'completed_at']:
                if job_data.get(key):
                    job_data[key] = datetime.fromisoformat(job_data[key])
            job = ResearchJob(**job_data)
            jobs[job.id] = job
            loaded_count += 1
        except Exception as e:
            logger.error(f"Failed to load job from {job_file}: {e}")
    logger.info(f"Loaded {loaded_count} jobs from disk")


# Load existing jobs on module import
_load_jobs_from_disk()

# WebSocket connections for progress updates
active_connections: Dict[str, WebSocket] = {}

# Thread pool for running research tasks
_executor = ThreadPoolExecutor(max_workers=4)


@router.get("/providers")
async def get_providers():
    """Get available AI providers and their status."""
    service = get_research_service()
    availability = service.get_available_providers()
    return {"providers": availability}


def _run_research_sync(job_id: str):
    """Synchronous function to run research in a background thread."""
    if job_id not in jobs:
        logger.error(f"Job {job_id} not found in jobs dict")
        return

    job = jobs[job_id]
    service = get_research_service()

    try:
        # Update status to running
        job.status = ResearchStatus.RUNNING
        job.updated_at = datetime.utcnow()

        # Initialize provider results as pending
        job.results = [
            ProviderResult(provider=p, status=ResearchStatus.PENDING)
            for p in job.providers
        ]
        job.progress = 10

        # Find relevant sources for this prompt
        prompt_keywords = get_prompt_keywords(job.prompt_id)
        relevant_sources = find_relevant_sources(prompt_keywords, limit=5)
        source_hints = format_sources_for_prompt(relevant_sources)

        # Store recommended source IDs on job for later reference
        job.recommended_sources = [s["id"] for s in relevant_sources]

        logger.info(f"Starting research for job {job_id}: {job.company} - {job.prompt_id}")
        logger.info(f"Found {len(relevant_sources)} relevant sources for prompt keywords: {prompt_keywords}")

        # Check if using DEEP mode with Open Deep Research (LangChain's multi-agent workflow)
        if job.mode == "deep" and OPEN_DEEP_RESEARCH_AVAILABLE:
            logger.info(f"Using Open Deep Research (LangChain) for job {job_id}")
            job.progress = 15

            # Get the prompt content
            prompt_content = service.engine.load_prompt(job.prompt_id)

            # Map our provider names to Open Deep Research provider names
            provider_mapping = {
                "openai": "openai",
                "anthropic": "anthropic",
                "google": "google",
                "xai": "xai",
                "perplexity": "perplexity",
            }

            # Create event loop for async execution
            loop = asyncio.new_event_loop()
            asyncio.set_event_loop(loop)
            try:
                # Run deep research for each provider in parallel
                odr_service = get_open_deep_research_service()

                # Check if service is available
                is_healthy = loop.run_until_complete(odr_service.check_health())
                if not is_healthy:
                    logger.warning("Open Deep Research server not available, falling back to basic research")
                    # Fall through to standard research flow
                else:
                    job.progress = 20

                    # Run parallel research with all selected providers
                    mapped_providers = [
                        provider_mapping.get(p, "openai")
                        for p in job.providers
                    ]

                    research_results = loop.run_until_complete(
                        odr_service.run_research_parallel(
                            query=prompt_content,
                            providers=mapped_providers,
                            company=job.company,
                            deep_mode=True,
                            # Uses DEEP_RESEARCH_TIMEOUT_SECONDS env var (default: 15 min)
                        )
                    )

                    job.progress = 80
                    logger.info(f"Open Deep Research completed for job {job_id}")

                    # Convert results to job format
                    all_citations = []
                    job.results = []

                    for orig_provider in job.providers:
                        mapped_provider = provider_mapping.get(orig_provider, "openai")
                        result = research_results.get(mapped_provider)

                        if result and result.status == "completed" and result.content:
                            job.results.append(ProviderResult(
                                provider=orig_provider,
                                status=ResearchStatus.COMPLETED,
                                content=result.content,
                                citations=result.citations or [],
                                model=f"open-deep-research:{result.model}",
                            ))
                            all_citations.extend(result.citations or [])
                        elif result and result.error:
                            job.results.append(ProviderResult(
                                provider=orig_provider,
                                status=ResearchStatus.FAILED,
                                error=result.error,
                            ))
                        else:
                            job.results.append(ProviderResult(
                                provider=orig_provider,
                                status=ResearchStatus.FAILED,
                                error="No result returned from Open Deep Research",
                            ))

                    # If multiple providers succeeded, synthesize results
                    successful_results = [r for r in job.results if r.status == ResearchStatus.COMPLETED]
                    if len(successful_results) >= 2:
                        job.progress = 85
                        logger.info(f"Synthesizing {len(successful_results)} Open Deep Research results for job {job_id}")
                        try:
                            synthesis_service = get_synthesis_service()
                            model_outputs = {
                                r.provider: r.content
                                for r in successful_results
                            }
                            synthesis_result = synthesis_service.synthesize(
                                company=job.company,
                                prompt_name=job.prompt_name,
                                model_outputs=model_outputs,
                            )
                            job.synthesized_report = SynthesizedReport(
                                content=synthesis_result.full_markdown,
                                models_used=synthesis_result.models_used,
                                high_confidence_findings=synthesis_result.high_confidence_findings,
                                areas_of_disagreement=synthesis_result.areas_of_disagreement,
                                unique_insights=synthesis_result.unique_insights,
                                generated_at=synthesis_result.generated_at,
                            )
                            logger.info(f"Synthesis complete for job {job_id}")
                        except Exception as synth_err:
                            logger.error(f"Synthesis failed for job {job_id}: {synth_err}")

                    # Match citations to curated sources and update hit counts
                    if all_citations:
                        matched_source_ids = match_citations_to_sources(all_citations)
                        if matched_source_ids:
                            job.cited_sources = matched_source_ids
                            update_source_hit_counts(matched_source_ids)
                            logger.info(f"Matched {len(matched_source_ids)} curated sources")

                    # Mark job as complete
                    job.status = ResearchStatus.COMPLETED
                    job.progress = 100
                    job.completed_at = datetime.utcnow()
                    job.updated_at = datetime.utcnow()
                    _save_job_to_disk(job)
                    logger.info(f"Open Deep Research job {job_id} completed successfully")
                    return  # Exit early, we're done

            except Exception as odr_error:
                logger.error(f"Open Deep Research failed, falling back to standard: {odr_error}")
                import traceback
                logger.error(traceback.format_exc())
                # Fall through to standard research flow
            finally:
                loop.close()

        # Check if using orchestrated mode with LangGraph
        if job.mode == "orchestrated" and LANGGRAPH_AVAILABLE:
            logger.info(f"Using LangGraph orchestration for job {job_id}")
            job.progress = 15

            # Get the prompt content
            prompt_content = service.engine.load_prompt(job.prompt_id)

            # Run orchestrated research (async) in a new event loop
            orchestrator = get_langgraph_orchestrator()

            # Create event loop for async execution
            loop = asyncio.new_event_loop()
            asyncio.set_event_loop(loop)
            try:
                orchestration_result = loop.run_until_complete(
                    orchestrator.run_deep_research(
                        company=job.company,
                        prompt=prompt_content,
                        prompt_id=job.prompt_id,
                        providers=job.providers,
                        mode="deep"  # Use deep research for each sub-topic
                    )
                )
            finally:
                loop.close()

            job.progress = 70
            logger.info(f"LangGraph orchestration completed for job {job_id}")

            # Process orchestration result
            if orchestration_result.get("status") == "completed" or orchestration_result.get("status") == "completed_with_errors":
                # Create a synthesized report from the orchestration output
                final_report = orchestration_result.get("final_report", "")
                citations = orchestration_result.get("citations", [])
                sub_topics = orchestration_result.get("sub_topics", [])

                # Store as synthesized report
                job.synthesized_report = SynthesizedReport(
                    content=final_report,
                    models_used=job.providers,
                    high_confidence_findings=[f"Sub-topic: {st.get('topic', 'N/A')}" for st in sub_topics[:5]],
                    areas_of_disagreement=[],
                    unique_insights={},
                    generated_at=datetime.utcnow(),
                )

                # Create a placeholder provider result for the orchestrated output
                job.results = [ProviderResult(
                    provider="orchestrated",
                    status=ResearchStatus.COMPLETED,
                    content=final_report,
                    citations=[{"url": c.get("url", ""), "title": c.get("title", "")} for c in citations],
                    model="langgraph-orchestrator",
                )]

                # Collect all citations for source matching
                all_citations = [{"url": c.get("url", ""), "title": c.get("title", "")} for c in citations]

            else:
                # Orchestration failed
                error_msg = orchestration_result.get("error", "Unknown orchestration error")
                job.results = [ProviderResult(
                    provider="orchestrated",
                    status=ResearchStatus.FAILED,
                    error=error_msg,
                )]
                all_citations = []

            logger.info(f"Orchestrated research completed for job {job_id}")

        else:
            # Standard research flow
            # Run the research synchronously
            # Note: source_hints can be appended to the prompt if the engine supports it
            results = service.engine.run_research(
                job.company,
                job.prompt_id,
                job.mode if job.mode != "orchestrated" else "deep",  # Fallback to deep if LangGraph unavailable
                job.providers
            )

            logger.info(f"Research completed for job {job_id}, processing results...")

            # Collect all citations for source matching
            all_citations = []

            # Update job with results
            job.results = []
            for provider, result in results.items():
                if result.error is None:
                    # Convert citations to the expected format
                    citations = [
                        {"url": c.get("url", ""), "title": c.get("title", "")}
                        for c in (result.citations or [])
                    ]
                    all_citations.extend(citations)

                    job.results.append(ProviderResult(
                        provider=provider,
                        status=ResearchStatus.COMPLETED,
                        content=result.text,
                        citations=citations,
                        model=result.model,
                    ))
                else:
                    job.results.append(ProviderResult(
                        provider=provider,
                        status=ResearchStatus.FAILED,
                        error=result.error,
                    ))

            # If mode is "combined", synthesize the results into a unified report
            if job.mode == "combined" and len(job.results) >= 2:
                job.progress = 80
                logger.info(f"Synthesizing {len(job.results)} model outputs for job {job_id}")
                try:
                    synthesis_service = get_synthesis_service()

                    # Collect successful model outputs
                    model_outputs = {}
                    for pr in job.results:
                        if pr.status == ResearchStatus.COMPLETED and pr.content:
                            model_outputs[pr.provider] = pr.content

                    if len(model_outputs) >= 2:
                        # Run synthesis
                        synthesis_result = synthesis_service.synthesize(
                            company=job.company,
                            prompt_name=job.prompt_name,
                            model_outputs=model_outputs,
                        )

                        # Store the synthesized report
                        job.synthesized_report = SynthesizedReport(
                            content=synthesis_result.full_markdown,
                            models_used=synthesis_result.models_used,
                            high_confidence_findings=synthesis_result.high_confidence_findings,
                            areas_of_disagreement=synthesis_result.areas_of_disagreement,
                            unique_insights=synthesis_result.unique_insights,
                            generated_at=synthesis_result.generated_at,
                        )
                        logger.info(f"Synthesis complete for job {job_id}")
                    else:
                        logger.warning(f"Not enough successful model outputs for synthesis in job {job_id}")
                except Exception as synth_err:
                    logger.error(f"Synthesis failed for job {job_id}: {synth_err}\n{traceback.format_exc()}")
                    # Don't fail the job, just skip synthesis

        # Match citations to curated sources and update hit counts
        if all_citations:
            matched_source_ids = match_citations_to_sources(all_citations)
            if matched_source_ids:
                job.cited_sources = matched_source_ids
                update_source_hit_counts(matched_source_ids)
                logger.info(f"Matched {len(matched_source_ids)} curated sources: {matched_source_ids}")

        # Save results to disk (only for non-orchestrated modes)
        if job.mode != "orchestrated" or not LANGGRAPH_AVAILABLE:
            save_info = service.save_results(results, job.company, job.prompt_name)
            job.result_file = save_info.get("output_dir")

        # Mark job as complete
        job.status = ResearchStatus.COMPLETED
        job.progress = 100
        job.completed_at = datetime.utcnow()
        job.updated_at = datetime.utcnow()

        logger.info(f"Research job {job_id} completed successfully")

        # Run knowledge base delta analysis (extracts entities & detects changes)
        try:
            # Combine all successful results for analysis
            combined_content = []
            for pr in job.results:
                if pr.content:
                    combined_content.append(f"## {pr.provider.upper()} Results\n\n{pr.content}")

            if combined_content:
                full_content = "\n\n---\n\n".join(combined_content)

                # Use the new knowledge base approach
                # This extracts entities, compares to the company's knowledge base,
                # and returns changes as delta findings
                delta_report = run_knowledge_base_delta(
                    company=job.company,
                    report_content=full_content,
                    job_id=job_id,
                    prompt_id=job.prompt_id,
                )
                if delta_report:
                    save_delta_report(delta_report)
                    job.delta_report_id = delta_report.id
                    logger.info(f"Knowledge base delta complete for job {job_id}: {len(delta_report.findings)} changes detected")
        except Exception as delta_err:
            logger.warning(f"Knowledge base delta analysis failed for job {job_id}: {delta_err}")

        # Save completed job to disk
        _save_job_to_disk(job)

    except Exception as e:
        logger.error(f"Research job {job_id} failed: {e}\n{traceback.format_exc()}")
        job.status = ResearchStatus.FAILED
        job.error = str(e)
        job.updated_at = datetime.utcnow()
        _save_job_to_disk(job)  # Save failed job to disk too


@router.post("", response_model=ResearchResponse)
async def start_research(request: ResearchRequest, background_tasks: BackgroundTasks):
    """Start a new research job"""
    job_id = str(uuid.uuid4())
    now = datetime.utcnow()

    # Create job record
    job = ResearchJob(
        id=job_id,
        company=request.company,
        prompt_id=request.prompt_id,
        prompt_name=request.prompt_id.replace("_", " ").title(),  # Simple name for now
        status=ResearchStatus.PENDING,
        providers=request.providers,
        mode=request.mode,
        created_at=now,
        updated_at=now,
    )

    jobs[job_id] = job
    _save_job_to_disk(job)  # Persist to disk

    # Start the research job in background thread
    _executor.submit(_run_research_sync, job_id)

    return ResearchResponse(
        job_id=job_id,
        status=ResearchStatus.PENDING,
        message=f"Research job created for {request.company}"
    )


@router.get("/{job_id}", response_model=ResearchJob)
async def get_job(job_id: str):
    """Get the status of a research job"""
    if job_id not in jobs:
        raise HTTPException(status_code=404, detail=f"Job not found: {job_id}")

    return jobs[job_id]


@router.get("", response_model=list[ResearchJob])
async def list_jobs(limit: int = 10):
    """List recent research jobs"""
    sorted_jobs = sorted(jobs.values(), key=lambda j: j.created_at, reverse=True)
    return sorted_jobs[:limit]


@router.websocket("/ws/{job_id}")
async def websocket_progress(websocket: WebSocket, job_id: str):
    """WebSocket endpoint for real-time job progress updates"""
    await websocket.accept()
    active_connections[job_id] = websocket

    try:
        while True:
            # Keep connection alive and send progress updates
            if job_id in jobs:
                job = jobs[job_id]
                progress = ResearchProgress(
                    job_id=job_id,
                    status=job.status,
                    current_step="Processing...",
                    progress_percent=50,  # Placeholder
                    timestamp=datetime.utcnow(),
                )
                await websocket.send_json(progress.model_dump(mode="json"))

                if job.status in [ResearchStatus.COMPLETED, ResearchStatus.FAILED]:
                    break

            await asyncio.sleep(1)

    except WebSocketDisconnect:
        pass
    finally:
        if job_id in active_connections:
            del active_connections[job_id]


def generate_professional_html(content: str, company: str, prompt_name: str, completed_at: str, provider: str) -> str:
    """Generate a professionally styled HTML document that looks like a PDF report."""
    html_body = markdown.markdown(content, extensions=['tables', 'fenced_code', 'toc'])

    return f"""<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="utf-8">
    <meta name="viewport" content="width=device-width, initial-scale=1">
    <title>{company} - {prompt_name} | Competitive Intelligence Report</title>
    <style>
        @import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700&family=Merriweather:wght@400;700&display=swap');

        :root {{
            --primary-color: #1e40af;
            --secondary-color: #3b82f6;
            --text-primary: #1f2937;
            --text-secondary: #6b7280;
            --border-color: #e5e7eb;
            --bg-light: #f9fafb;
        }}

        * {{
            margin: 0;
            padding: 0;
            box-sizing: border-box;
        }}

        @page {{
            size: A4;
            margin: 2cm;
        }}

        body {{
            font-family: 'Inter', -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif;
            font-size: 11pt;
            line-height: 1.7;
            color: var(--text-primary);
            background: white;
            max-width: 850px;
            margin: 0 auto;
            padding: 40px;
        }}

        /* Cover/Header Section */
        .report-header {{
            border-bottom: 3px solid var(--primary-color);
            padding-bottom: 30px;
            margin-bottom: 40px;
        }}

        .report-header .logo {{
            font-size: 14pt;
            font-weight: 700;
            color: var(--primary-color);
            letter-spacing: -0.5px;
            margin-bottom: 30px;
        }}

        .report-header h1 {{
            font-family: 'Merriweather', Georgia, serif;
            font-size: 28pt;
            font-weight: 700;
            color: var(--text-primary);
            margin-bottom: 10px;
            line-height: 1.2;
        }}

        .report-header .subtitle {{
            font-size: 14pt;
            color: var(--secondary-color);
            font-weight: 500;
            margin-bottom: 20px;
        }}

        .report-meta {{
            display: flex;
            gap: 30px;
            font-size: 10pt;
            color: var(--text-secondary);
        }}

        .report-meta div {{
            display: flex;
            align-items: center;
            gap: 6px;
        }}

        .report-meta strong {{
            color: var(--text-primary);
            font-weight: 600;
        }}

        /* Content Styling */
        .report-content {{
            page-break-inside: auto;
        }}

        h1 {{
            font-family: 'Merriweather', Georgia, serif;
            font-size: 20pt;
            font-weight: 700;
            color: var(--primary-color);
            margin: 35px 0 20px;
            padding-bottom: 10px;
            border-bottom: 2px solid var(--border-color);
            page-break-after: avoid;
        }}

        h2 {{
            font-family: 'Merriweather', Georgia, serif;
            font-size: 16pt;
            font-weight: 700;
            color: var(--text-primary);
            margin: 30px 0 15px;
            page-break-after: avoid;
        }}

        h3 {{
            font-size: 13pt;
            font-weight: 600;
            color: var(--text-primary);
            margin: 25px 0 12px;
            page-break-after: avoid;
        }}

        h4 {{
            font-size: 11pt;
            font-weight: 600;
            color: var(--text-secondary);
            margin: 20px 0 10px;
            text-transform: uppercase;
            letter-spacing: 0.5px;
        }}

        p {{
            margin-bottom: 14px;
            text-align: justify;
            orphans: 3;
            widows: 3;
        }}

        ul, ol {{
            margin: 15px 0 15px 25px;
        }}

        li {{
            margin-bottom: 8px;
        }}

        li::marker {{
            color: var(--secondary-color);
        }}

        /* Tables */
        table {{
            width: 100%;
            border-collapse: collapse;
            margin: 20px 0;
            font-size: 10pt;
            page-break-inside: avoid;
        }}

        thead {{
            background: var(--primary-color);
            color: white;
        }}

        th {{
            padding: 12px 15px;
            text-align: left;
            font-weight: 600;
            text-transform: uppercase;
            font-size: 9pt;
            letter-spacing: 0.5px;
        }}

        td {{
            padding: 12px 15px;
            border-bottom: 1px solid var(--border-color);
        }}

        tbody tr:nth-child(even) {{
            background: var(--bg-light);
        }}

        tbody tr:hover {{
            background: #f3f4f6;
        }}

        /* Blockquotes */
        blockquote {{
            border-left: 4px solid var(--secondary-color);
            background: var(--bg-light);
            margin: 20px 0;
            padding: 15px 20px;
            font-style: italic;
            color: var(--text-secondary);
        }}

        blockquote p:last-child {{
            margin-bottom: 0;
        }}

        /* Code */
        code {{
            font-family: 'SF Mono', Monaco, 'Courier New', monospace;
            background: var(--bg-light);
            padding: 2px 6px;
            border-radius: 4px;
            font-size: 10pt;
        }}

        pre {{
            background: #1f2937;
            color: #f9fafb;
            padding: 20px;
            border-radius: 8px;
            overflow-x: auto;
            margin: 20px 0;
        }}

        pre code {{
            background: none;
            padding: 0;
            color: inherit;
        }}

        /* Links */
        a {{
            color: var(--secondary-color);
            text-decoration: none;
        }}

        a:hover {{
            text-decoration: underline;
        }}

        /* Horizontal Rule */
        hr {{
            border: none;
            border-top: 2px solid var(--border-color);
            margin: 30px 0;
        }}

        /* Strong/Bold */
        strong {{
            font-weight: 600;
            color: var(--text-primary);
        }}

        /* Emphasis */
        em {{
            font-style: italic;
        }}

        /* Footer */
        .report-footer {{
            margin-top: 50px;
            padding-top: 20px;
            border-top: 2px solid var(--border-color);
            font-size: 9pt;
            color: var(--text-secondary);
            text-align: center;
        }}

        /* Print Styles */
        @media print {{
            body {{
                padding: 0;
                max-width: none;
            }}

            .report-header {{
                page-break-after: avoid;
            }}

            h1, h2, h3 {{
                page-break-after: avoid;
            }}

            table, figure {{
                page-break-inside: avoid;
            }}
        }}
    </style>
</head>
<body>
    <header class="report-header">
        <div class="logo">RED 6 COMPETITIVE INTELLIGENCE</div>
        <h1>{company}</h1>
        <div class="subtitle">{prompt_name}</div>
        <div class="report-meta">
            <div><strong>Generated:</strong> {completed_at}</div>
            <div><strong>Source:</strong> {provider.upper()}</div>
            <div><strong>Classification:</strong> Internal Use Only</div>
        </div>
    </header>

    <main class="report-content">
        {html_body}
    </main>

    <footer class="report-footer">
        <p>This report was generated by Red 6 Competitive Intelligence Platform</p>
        <p>Confidential - For Internal Use Only</p>
    </footer>
</body>
</html>"""


@router.get("/{job_id}/export/{provider}")
async def export_result(
    job_id: str,
    provider: str,
    format: Literal["md", "html", "txt", "pdf", "docx"] = "md"
):
    """Export a specific provider's result in various formats"""
    # Handle "combined" separately - redirect to the combined export endpoint
    if provider == "combined":
        return await export_combined_result(job_id, format)

    if job_id not in jobs:
        raise HTTPException(status_code=404, detail=f"Job not found: {job_id}")

    job = jobs[job_id]

    if job.status != ResearchStatus.COMPLETED:
        raise HTTPException(status_code=400, detail="Job not completed yet")

    # Find the provider result
    result = next((r for r in job.results if r.provider == provider), None)
    if not result:
        raise HTTPException(status_code=404, detail=f"Provider {provider} not found in results")

    if not result.content:
        raise HTTPException(status_code=400, detail="No content available for export")

    content = result.content
    filename = f"{job.company}_{job.prompt_id}_{provider}"
    completed_at = job.completed_at.strftime("%B %d, %Y at %I:%M %p") if job.completed_at else "N/A"

    if format == "md":
        return Response(
            content=content,
            media_type="text/markdown",
            headers={"Content-Disposition": f'attachment; filename="{filename}.md"'}
        )
    elif format == "html":
        html_content = generate_professional_html(content, job.company, job.prompt_name, completed_at, provider)
        return Response(
            content=html_content,
            media_type="text/html",
            headers={"Content-Disposition": f'attachment; filename="{filename}.html"'}
        )
    elif format == "pdf":
        try:
            from weasyprint import HTML
            html_content = generate_professional_html(content, job.company, job.prompt_name, completed_at, provider)
            pdf_buffer = io.BytesIO()
            HTML(string=html_content).write_pdf(pdf_buffer)
            pdf_buffer.seek(0)
            return Response(
                content=pdf_buffer.getvalue(),
                media_type="application/pdf",
                headers={"Content-Disposition": f'attachment; filename="{filename}.pdf"'}
            )
        except ImportError:
            raise HTTPException(status_code=500, detail="PDF export not available. Install weasyprint.")
        except Exception as e:
            logger.error(f"PDF generation failed: {e}")
            raise HTTPException(status_code=500, detail=f"PDF generation failed: {str(e)}")
    elif format == "docx":
        try:
            from docx import Document
            from docx.shared import Inches, Pt, RGBColor
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            from docx.enum.style import WD_STYLE_TYPE

            doc = Document()

            # Set document margins
            for section in doc.sections:
                section.top_margin = Inches(1)
                section.bottom_margin = Inches(1)
                section.left_margin = Inches(1.25)
                section.right_margin = Inches(1.25)

            # Add header
            header_para = doc.add_paragraph()
            header_run = header_para.add_run("RED 6 COMPETITIVE INTELLIGENCE")
            header_run.bold = True
            header_run.font.size = Pt(12)
            header_run.font.color.rgb = RGBColor(30, 64, 175)

            # Add title
            title = doc.add_heading(job.company, 0)
            title.alignment = WD_ALIGN_PARAGRAPH.LEFT

            # Add subtitle
            subtitle = doc.add_paragraph()
            subtitle_run = subtitle.add_run(job.prompt_name)
            subtitle_run.font.size = Pt(14)
            subtitle_run.font.color.rgb = RGBColor(59, 130, 246)

            # Add metadata
            meta = doc.add_paragraph()
            meta.add_run(f"Generated: {completed_at}  |  Source: {provider.upper()}  |  Classification: Internal Use Only")
            meta.runs[0].font.size = Pt(10)
            meta.runs[0].font.color.rgb = RGBColor(107, 114, 128)

            doc.add_paragraph()  # Spacer

            # Parse and add content
            lines = content.split('\n')
            for line in lines:
                line = line.strip()
                if not line:
                    doc.add_paragraph()
                elif line.startswith('# '):
                    doc.add_heading(line[2:], 1)
                elif line.startswith('## '):
                    doc.add_heading(line[3:], 2)
                elif line.startswith('### '):
                    doc.add_heading(line[4:], 3)
                elif line.startswith('#### '):
                    doc.add_heading(line[5:], 4)
                elif line.startswith('- ') or line.startswith('* '):
                    para = doc.add_paragraph(line[2:], style='List Bullet')
                elif re.match(r'^\d+\.\s', line):
                    para = doc.add_paragraph(re.sub(r'^\d+\.\s', '', line), style='List Number')
                elif line.startswith('> '):
                    para = doc.add_paragraph()
                    para.paragraph_format.left_indent = Inches(0.5)
                    run = para.add_run(line[2:])
                    run.italic = True
                    run.font.color.rgb = RGBColor(107, 114, 128)
                else:
                    # Handle bold and italic in regular text
                    para = doc.add_paragraph()
                    # Simple bold/italic handling
                    text = line
                    text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)  # Remove bold markers
                    text = re.sub(r'\*(.+?)\*', r'\1', text)  # Remove italic markers
                    para.add_run(text)

            # Add footer
            doc.add_paragraph()
            footer = doc.add_paragraph()
            footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
            footer_run = footer.add_run("Generated by Red 6 Competitive Intelligence Platform - Confidential")
            footer_run.font.size = Pt(9)
            footer_run.font.color.rgb = RGBColor(156, 163, 175)

            docx_buffer = io.BytesIO()
            doc.save(docx_buffer)
            docx_buffer.seek(0)

            return Response(
                content=docx_buffer.getvalue(),
                media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                headers={"Content-Disposition": f'attachment; filename="{filename}.docx"'}
            )
        except ImportError:
            raise HTTPException(status_code=500, detail="DOCX export not available. Install python-docx.")
        except Exception as e:
            logger.error(f"DOCX generation failed: {e}")
            raise HTTPException(status_code=500, detail=f"DOCX generation failed: {str(e)}")
    else:  # txt
        return Response(
            content=content,
            media_type="text/plain",
            headers={"Content-Disposition": f'attachment; filename="{filename}.txt"'}
        )


@router.get("/{job_id}/export/combined")
async def export_combined_result(
    job_id: str,
    format: Literal["md", "html", "txt", "pdf", "docx"] = "md"
):
    """Export the synthesized combined report"""
    if job_id not in jobs:
        raise HTTPException(status_code=404, detail=f"Job not found: {job_id}")

    job = jobs[job_id]

    if job.status != ResearchStatus.COMPLETED:
        raise HTTPException(status_code=400, detail="Job not completed yet")

    if not job.synthesized_report:
        raise HTTPException(status_code=404, detail="No synthesized report available for this job")

    content = job.synthesized_report.content
    filename = f"{job.company}_{job.prompt_id}_combined"
    completed_at = job.completed_at.strftime("%B %d, %Y at %I:%M %p") if job.completed_at else "N/A"
    models_str = ", ".join(job.synthesized_report.models_used)

    if format == "md":
        return Response(
            content=content,
            media_type="text/markdown",
            headers={"Content-Disposition": f'attachment; filename="{filename}.md"'}
        )
    elif format == "html":
        html_content = generate_professional_html(content, job.company, job.prompt_name, completed_at, f"Combined ({models_str})")
        return Response(
            content=html_content,
            media_type="text/html",
            headers={"Content-Disposition": f'attachment; filename="{filename}.html"'}
        )
    elif format == "pdf":
        try:
            from weasyprint import HTML
            html_content = generate_professional_html(content, job.company, job.prompt_name, completed_at, f"Combined ({models_str})")
            pdf_buffer = io.BytesIO()
            HTML(string=html_content).write_pdf(pdf_buffer)
            pdf_buffer.seek(0)
            return Response(
                content=pdf_buffer.getvalue(),
                media_type="application/pdf",
                headers={"Content-Disposition": f'attachment; filename="{filename}.pdf"'}
            )
        except ImportError:
            raise HTTPException(status_code=500, detail="PDF export not available. Install weasyprint.")
        except Exception as e:
            logger.error(f"PDF generation failed: {e}")
            raise HTTPException(status_code=500, detail=f"PDF generation failed: {str(e)}")
    elif format == "docx":
        try:
            from docx import Document
            from docx.shared import Inches, Pt, RGBColor
            from docx.enum.text import WD_ALIGN_PARAGRAPH

            doc = Document()

            for section in doc.sections:
                section.top_margin = Inches(1)
                section.bottom_margin = Inches(1)
                section.left_margin = Inches(1.25)
                section.right_margin = Inches(1.25)

            header_para = doc.add_paragraph()
            header_run = header_para.add_run("RED 6 COMPETITIVE INTELLIGENCE - COMBINED ANALYSIS")
            header_run.bold = True
            header_run.font.size = Pt(12)
            header_run.font.color.rgb = RGBColor(30, 64, 175)

            title = doc.add_heading(job.company, 0)
            title.alignment = WD_ALIGN_PARAGRAPH.LEFT

            subtitle = doc.add_paragraph()
            subtitle_run = subtitle.add_run(job.prompt_name)
            subtitle_run.font.size = Pt(14)
            subtitle_run.font.color.rgb = RGBColor(59, 130, 246)

            meta = doc.add_paragraph()
            meta.add_run(f"Generated: {completed_at}  |  Models: {models_str}  |  Classification: Internal Use Only")
            meta.runs[0].font.size = Pt(10)
            meta.runs[0].font.color.rgb = RGBColor(107, 114, 128)

            doc.add_paragraph()

            lines = content.split('\n')
            for line in lines:
                line = line.strip()
                if not line:
                    doc.add_paragraph()
                elif line.startswith('# '):
                    doc.add_heading(line[2:], 1)
                elif line.startswith('## '):
                    doc.add_heading(line[3:], 2)
                elif line.startswith('### '):
                    doc.add_heading(line[4:], 3)
                elif line.startswith('#### '):
                    doc.add_heading(line[5:], 4)
                elif line.startswith('- ') or line.startswith('* '):
                    doc.add_paragraph(line[2:], style='List Bullet')
                elif re.match(r'^\d+\.\s', line):
                    doc.add_paragraph(re.sub(r'^\d+\.\s', '', line), style='List Number')
                elif line.startswith('> '):
                    para = doc.add_paragraph()
                    para.paragraph_format.left_indent = Inches(0.5)
                    run = para.add_run(line[2:])
                    run.italic = True
                    run.font.color.rgb = RGBColor(107, 114, 128)
                else:
                    para = doc.add_paragraph()
                    text = line
                    text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
                    text = re.sub(r'\*(.+?)\*', r'\1', text)
                    para.add_run(text)

            doc.add_paragraph()
            footer = doc.add_paragraph()
            footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
            footer_run = footer.add_run("Generated by Red 6 Competitive Intelligence Platform - Confidential")
            footer_run.font.size = Pt(9)
            footer_run.font.color.rgb = RGBColor(156, 163, 175)

            docx_buffer = io.BytesIO()
            doc.save(docx_buffer)
            docx_buffer.seek(0)

            return Response(
                content=docx_buffer.getvalue(),
                media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                headers={"Content-Disposition": f'attachment; filename="{filename}.docx"'}
            )
        except ImportError:
            raise HTTPException(status_code=500, detail="DOCX export not available. Install python-docx.")
        except Exception as e:
            logger.error(f"DOCX generation failed: {e}")
            raise HTTPException(status_code=500, detail=f"DOCX generation failed: {str(e)}")
    else:  # txt
        return Response(
            content=content,
            media_type="text/plain",
            headers={"Content-Disposition": f'attachment; filename="{filename}.txt"'}
        )


@router.post("/{job_id}/presentation")
async def generate_presentation(
    job_id: str,
    provider: str = "combined",
    llm_provider: str = "google"
):
    """Generate slide content as text that can be copied into a branded slide deck."""
    from services.presentation_service import get_presentation_service
    from models.research import ResearchStatus

    job = await get_job(job_id)
    if not job:
        raise HTTPException(status_code=404, detail="Job not found")

    if job.status != ResearchStatus.COMPLETED:
        raise HTTPException(status_code=400, detail="Job is not completed")

    company = job.company or "Company"
    prompt_name = job.prompt_name or "Research"
    content = None
    citations = []

    if provider == "combined":
        # Use combined/synthesized content
        if not job.synthesized_report:
            raise HTTPException(status_code=400, detail="No combined result available")
        content = job.synthesized_report.content
        # SynthesizedReport doesn't have citations, collect from all providers
        for result in job.results:
            if result.citations:
                citations.extend(result.citations)
    else:
        # Use specific provider result - find in results list
        provider_result = None
        for result in job.results:
            if result.provider == provider:
                provider_result = result
                break
        if not provider_result:
            raise HTTPException(status_code=400, detail=f"No result found for provider: {provider}")
        content = provider_result.content or ""
        citations = provider_result.citations or []

    if not content:
        raise HTTPException(status_code=400, detail="No content available for presentation generation")

    try:
        presentation_service = get_presentation_service()
        # Convert Citation objects to dicts for the service
        citations_dicts = [{"url": c.url, "title": c.title} for c in citations] if citations else []
        slide_text = presentation_service.generate_slide_text(
            content=content,
            company=company,
            prompt_name=prompt_name,
            citations=citations_dicts
        )

        filename = f"{company}_{prompt_name}_slides".replace(" ", "_")

        return Response(
            content=slide_text,
            media_type="text/plain; charset=utf-8",
            headers={"Content-Disposition": f'attachment; filename="{filename}.txt"'}
        )
    except Exception as e:
        logger.error(f"Slide content generation failed: {e}")
        raise HTTPException(status_code=500, detail=f"Slide content generation failed: {str(e)}")
